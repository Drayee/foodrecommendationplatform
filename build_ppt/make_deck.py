#!/usr/bin/env python3
"""生成《寻味 · 美食推荐平台》课程项目汇报 PPT（16:9）。

用法：python build_ppt/make_deck.py
输出：docs/寻味美食推荐平台-第十组-项目汇报.pptx
"""
from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets"
OUT = ROOT / "docs" / "寻味美食推荐平台-第十组-项目汇报.pptx"

# ---------------------------------------------------------------- 设计系统
FONT = "微软雅黑"
FONT_NUM = "Segoe UI Semibold"

INK = "1B2030"
INK_SOFT = "333A4D"
GRAY = "6E7686"
GRAY_L = "9AA1AE"
ORANGE = "F97316"
ORANGE_D = "D95F0A"
ORANGE_L = "FFE9D8"
CARD = "FFF7F0"
CARD_ALT = "F5F7FB"
TEAL = "0F9B8E"
TEAL_L = "DCF1EE"
LINE = "EBE0D6"
WHITE = "FFFFFF"

W, H = 13.333, 7.5


def I(v: float) -> Inches:
    return Inches(v)


def rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def style_run(run, size=12, bold=False, color=INK, font=FONT, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def tx(slide, x, y, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, spacing=1.28, space_after=0, font=FONT, margins=(0, 0, 0, 0)):
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = I(margins[0])
    tf.margin_right = I(margins[1])
    tf.margin_top = I(margins[2])
    tf.margin_bottom = I(margins[3])
    lines = str(text).split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        run = p.add_run()
        run.text = line
        style_run(run, size=size, bold=bold, color=color, font=font)
    return box


def rich(slide, x, y, w, h, segments, size=11, spacing=1.3, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0):
    """segments: list of list[(text, color, bold, size|None)]，每内层 list 是一段。"""
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, I(0))
    for idx, para in enumerate(segments):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        for seg in para:
            text, color, bold = seg[0], seg[1], seg[2]
            s = seg[3] if len(seg) > 3 else size
            run = p.add_run()
            run.text = text
            style_run(run, size=s, bold=bold, color=color)
    return box


def shadow_on(shape, blur=9, dist=4, alpha=11000):
    spPr = shape._element.spPr
    # shadow.inherit = False 已经写入一个空的 a:effectLst，先清掉避免重复元素
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}" '
        'dir="5400000" rotWithShape="0">'
        f'<a:srgbClr val="40210A"><a:alpha val="{alpha}"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst>"
    )
    spPr.append(parse_xml(xml))


def rect(slide, x, y, w, h, fill, line=None, radius=None, shadow=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if shadow:
        shadow_on(shp)
    shp.text_frame.word_wrap = True
    return shp


def pill(slide, x, y, w, h, text, fill=ORANGE_L, color=ORANGE_D, size=10, bold=True, line=None):
    shp = rect(slide, x, y, w, h, fill, line=line, radius=0.5)
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = I(0.06)
    tf.margin_right = I(0.06)
    tf.margin_top = I(0)
    tf.margin_bottom = I(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    style_run(run, size=size, bold=bold, color=color)
    return shp


def circle_num(slide, x, y, d, text, fill=ORANGE, color=WHITE, size=12):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    style_run(run, size=size, bold=True, color=color)
    return shp


def arrow(slide, x, y, w=0.28, h=0.17, color=ORANGE_L):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, I(x), I(y), I(w), I(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


# ---------------------------------------------------------------- 页面骨架
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def footer(slide, page, total=16):
    rect(slide, 0.62, 6.93, 12.1, 0.012, LINE, shape=MSO_SHAPE.RECTANGLE)
    tx(slide, 0.62, 7.02, 8.0, 0.3, "寻味 · 美食推荐平台　|　第十组项目汇报", size=9, color=GRAY_L)
    tx(slide, 10.6, 7.02, 2.12, 0.3, f"{page:02d} / {total:02d}", size=9, color=GRAY_L,
       align=PP_ALIGN.RIGHT, font=FONT_NUM)


def head(slide, kicker, title, sub=None, page=1):
    """内容页页眉。"""
    rect(slide, 0.62, 0.56, 0.075, 0.62, ORANGE, shape=MSO_SHAPE.RECTANGLE)
    tx(slide, 0.85, 0.48, 9.4, 0.66, title, size=25, bold=True, color=INK, spacing=1.05)
    tx(slide, 10.2, 0.6, 2.52, 0.3, kicker, size=9.5, bold=True, color=ORANGE_D, align=PP_ALIGN.RIGHT)
    if sub:
        tx(slide, 0.87, 1.06, 11.6, 0.34, sub, size=11, color=GRAY)
    footer(slide, page)


def section_label(slide, x, y, text, color=GRAY_L, size=9.5):
    tx(slide, x, y, 6.0, 0.26, text, size=size, bold=True, color=color)


# ---------------------------------------------------------------- 卡片构件
def feature_card(slide, x, y, w, h, num, title, body, accent=ORANGE, fill=CARD,
                 num_fill=None, title_color=INK, body_color=GRAY, body_size=10.5, radius=0.07):
    rect(slide, x, y, w, h, fill, radius=radius, shadow=True)
    circle_num(slide, x + 0.28, y + 0.26, 0.34, num, fill=num_fill or accent, size=10.5)
    tx(slide, x + 0.72, y + 0.24, w - 0.95, 0.34, title, size=13, bold=True, color=title_color)
    tx(slide, x + 0.3, y + 0.72, w - 0.6, h - 0.86, body, size=body_size, color=body_color, spacing=1.34)


def stat_card(slide, x, y, w, h, value, unit, label, note="", accent=ORANGE, fill=WHITE):
    rect(slide, x, y, w, h, fill, line=LINE, radius=0.08, shadow=True)
    rect(slide, x, y, w, 0.055, accent, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
    rich(slide, x + 0.2, y + 0.26, w - 0.4, 0.55,
         [[(value, accent, True, 27), (unit, GRAY, False, 11)]], spacing=1.0)
    tx(slide, x + 0.2, y + 0.84, w - 0.4, 0.28, label, size=11, bold=True, color=INK_SOFT)
    if note:
        tx(slide, x + 0.2, y + 1.12, w - 0.4, 0.3, note, size=9, color=GRAY_L)


def table(slide, x, y, widths, rows, row_h=0.62, header_h=0.42,
          head_fill=INK, head_color=WHITE, head_size=10.5, body_size=10,
          stripe=("FFFFFF", "FBF7F3"), radius=0.04, col_align=None, col_colors=None):
    """rows[0] 为表头。"""
    col_align = col_align or [PP_ALIGN.LEFT] * len(widths)
    cur = y
    total_w = sum(widths)
    rect(slide, x, cur, total_w, header_h, head_fill, radius=radius, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    cx = x
    for i, cell in enumerate(rows[0]):
        tx(slide, cx + 0.18, cur + (header_h - 0.26) / 2, widths[i] - 0.3, 0.26, cell,
           size=head_size, bold=True, color=head_color, align=col_align[i])
        cx += widths[i]
    cur += header_h + 0.06
    for r, row in enumerate(rows[1:]):
        fill = stripe[r % 2]
        rect(slide, x, cur, total_w, row_h, fill, line=LINE, radius=radius * 0.7)
        cx = x
        for i, cell in enumerate(row):
            color = INK_SOFT
            if col_colors and col_colors[i]:
                color = col_colors[i]
            bold = i == 0
            if bold:
                color = INK
            tx(slide, cx + 0.18, cur + 0.08, widths[i] - 0.3, row_h - 0.14, cell,
               size=body_size, bold=bold, color=color, align=col_align[i], spacing=1.18)
            cx += widths[i]
        cur += row_h + 0.055
    return cur


# ---------------------------------------------------------------- 封面
def slide_cover(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, INK, shape=MSO_SHAPE.RECTANGLE)
    # 装饰
    c1 = s.shapes.add_shape(MSO_SHAPE.OVAL, I(10.2), I(4.5), I(4.6), I(4.6))
    c1.fill.solid()
    c1.fill.fore_color.rgb = rgb("242A3D")
    c1.line.fill.background()
    c1.shadow.inherit = False
    c2 = s.shapes.add_shape(MSO_SHAPE.OVAL, I(11.6), I(-1.4), I(3.4), I(3.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = rgb(ORANGE)
    c2.line.fill.background()
    c2.shadow.inherit = False
    c3 = s.shapes.add_shape(MSO_SHAPE.OVAL, I(-1.1), I(5.4), I(2.9), I(2.9))
    c3.fill.solid()
    c3.fill.fore_color.rgb = rgb("242A3D")
    c3.line.fill.background()
    c3.shadow.inherit = False

    pill(s, 0.95, 0.92, 2.35, 0.4, "第十组 · 课程项目汇报", fill="2A3148", color="FFB37A", size=10.5)
    tx(s, 0.95, 1.62, 10.4, 1.0, "寻味 · 美食推荐平台", size=46, bold=True, color=WHITE)
    rect(s, 0.98, 2.78, 1.18, 0.075, ORANGE, shape=MSO_SHAPE.RECTANGLE)
    tx(s, 0.95, 3.06, 10.6, 1.0,
       "菜谱语义检索　×　个性化混合推荐　×　附近美食发现", size=17, color="FFD9BC")
    tx(s, 0.95, 3.62, 10.2, 0.9,
       "一个 Django 单项目，两条主线：在家做，出门吃。\n外部能力全部可降级 —— 没有 API Key、没有大模型、断网也能完成完整演示。",
       size=11.5, color="A8B0C2", spacing=1.5)

    tags = ["Django 6.1", "PostgreSQL 18", "RAG + BM25/RRF", "item-CF + MMR", "高德开放平台", "Ollama qwen3:14b"]
    x = 0.95
    for t in tags:
        wdt = 0.42 + len(t) * 0.098
        pill(s, x, 4.62, wdt, 0.34, t, fill="252C41", color="C9D1E0", size=9.5, bold=False)
        x += wdt + 0.14

    rect(s, 0.95, 5.42, 8.6, 0.012, "3A4257", shape=MSO_SHAPE.RECTANGLE)
    rich(s, 0.95, 5.62, 5.2, 0.9,
         [[("项目组长　", GRAY_L, False, 10), ("孔哲昊", WHITE, True, 12)]], spacing=1.2)
    rich(s, 0.95, 5.99, 9.6, 0.9,
         [[("小组成员　", GRAY_L, False, 10),
           ("彭冰钦　姚一飞　贺江涛　倪永硕　易子涵", WHITE, True, 12)]], spacing=1.2)
    tx(s, 0.95, 6.6, 6.0, 0.4, "2026 年 10 月", size=10, color="7C8598")
    return s


# ---------------------------------------------------------------- 目录
def slide_toc(prs, page):
    s = blank(prs)
    rect(s, 0, 0, W, H, "FDF9F5", shape=MSO_SHAPE.RECTANGLE)
    rect(s, 0, 0, 0.34, H, ORANGE, shape=MSO_SHAPE.RECTANGLE)
    tx(s, 0.95, 0.82, 6.0, 0.5, "CONTENTS", size=10.5, bold=True, color=ORANGE_D)
    tx(s, 0.95, 1.14, 6.0, 0.7, "汇报目录", size=32, bold=True, color=INK)
    tx(s, 0.95, 1.9, 5.6, 0.5, "从问题出发，讲清我们做了什么、怎么做的、做到什么程度。", size=11, color=GRAY)

    items = [
        ("01", "项目背景", "三个真实痛点：决策成本、推荐失焦、场景割裂"),
        ("02", "目标与范围", "四条可验收的目标与量化口径"),
        ("03", "系统架构", "六层职责划分与数据流"),
        ("04", "核心功能", "找菜主线 / 找店主线 双线并进"),
        ("05", "算法设计", "RAG 混合检索 与 四路混合推荐"),
        ("06", "工程质量", "降级矩阵、测试体系、契约驱动协作"),
        ("07", "数据与演示", "规模看板与三大界面实拍"),
        ("08", "团队与总结", "分工、难点对策、不足与展望"),
    ]
    x0, y0 = 0.95, 2.55
    for i, (num, title, desc) in enumerate(items):
        col = i % 2
        row = i // 2
        x = x0 + col * 5.85
        y = y0 + row * 1.06
        rect(s, x, y, 5.5, 0.86, WHITE, line=LINE, radius=0.09, shadow=True)
        tx(s, x + 0.24, y + 0.2, 0.7, 0.4, num, size=16, bold=True, color=ORANGE, font=FONT_NUM)
        tx(s, x + 0.86, y + 0.15, 1.9, 0.3, title, size=12.5, bold=True, color=INK)
        tx(s, x + 0.86, y + 0.46, 4.3, 0.3, desc, size=9, color=GRAY)
    footer(s, page)


# ---------------------------------------------------------------- 01 背景
def slide_background(prs, page):
    s = blank(prs)
    head(s, "01 / BACKGROUND", "项目背景：一顿饭的三个卡点",
         "我们要解决的不是「有没有菜谱」，而是「决定吃什么」这件事本身太贵。", page)

    cards = [
        ("01", "决策成本：站在冰箱前的一小时",
         "菜谱平台以目录和关键词为核心，用户检索「20 分钟、冰箱现有的食材、微辣」这类复合需求时，\n只能自己一页页翻。真正稀缺的不是菜谱，而是能直接落到方案上的建议。"),
        ("02", "推荐失焦：千人一面的榜单",
         "大量产品只按热度或最新排序，忽略口味偏好、忌口、可用时长与人均预算。\n结果是大盘数据好看的榜单，对「我」这个人却几乎无效，越刷越同质。"),
        ("03", "场景割裂：会做菜，却不知去哪吃",
         "菜谱社区不关心附近哪家店做得好，探店点评又给不出这道菜该怎么做。\n一次「吃」的决策被拆到两个 App 里，还要自己拼。"),
    ]
    y = 1.62
    for num, title, body in cards:
        rect(s, 0.62, y, 7.45, 1.5, CARD, radius=0.06, shadow=True)
        rect(s, 0.62, y + 0.28, 0.055, 0.94, ORANGE_L, shape=MSO_SHAPE.RECTANGLE)
        tx(s, 0.86, y + 0.2, 0.6, 0.36, num, size=15, bold=True, color=ORANGE, font=FONT_NUM)
        tx(s, 1.42, y + 0.22, 6.4, 0.34, title, size=13, bold=True, color=INK)
        tx(s, 1.42, y + 0.62, 6.4, 0.8, body, size=10, color=GRAY, spacing=1.35)
        y += 1.66

    # 右侧答案
    rect(s, 8.4, 1.62, 4.3, 4.68, INK, radius=0.05, shadow=True)
    tx(s, 8.72, 1.92, 3.7, 0.34, "我们的答案", size=10, bold=True, color="FFB37A")
    tx(s, 8.72, 2.2, 3.7, 0.9, "一个平台\n两条主线", size=23, bold=True, color=WHITE, spacing=1.18)
    rect(s, 8.74, 3.16, 0.75, 0.05, ORANGE, shape=MSO_SHAPE.RECTANGLE)

    blocks = [
        ("在家做", "菜谱库 + 语义检索 + 混合推荐 + AI 助手"),
        ("出门吃", "附近 POI + 五路排序 + 地图导航 + 打卡"),
    ]
    y = 3.42
    for title, body in blocks:
        tx(s, 8.74, y, 1.5, 0.3, title, size=12, bold=True, color="FFD9BC")
        tx(s, 8.74, y + 0.32, 3.66, 0.6, body, size=9.5, color="AEB6C7", spacing=1.35)
        y += 1.02

    tx(s, 8.74, 5.5, 3.66, 0.7,
       "统一账号体系与行为埋点：\n两条主线共享同一份口味画像。", size=9.5, color="8F98AC", spacing=1.4)


# ---------------------------------------------------------------- 02 目标
def slide_goals(prs, page):
    s = blank(prs)
    head(s, "02 / GOALS", "目标与范围：四条可验收的目标",
         "每条目标都对应一个可量化、可演示的验收口径，而不是「做一个美食平台」这种空话。", page)

    goals = [
        ("A", "找得快", "一句话完成检索",
         "自然语言输入直接命中意图：「想吃点辣的」「20 分钟快手菜」「减脂餐」。语义检索 + 标签召回双通道，\n不依赖用户学会分类目录。", ORANGE),
        ("B", "推得准", "四路加权 + 去同质化",
         "内容画像 0.45 + 协同过滤 0.30 + 热度 0.15 + 新鲜度 0.10；MMR 控制一屏内的菜系多样性，\n新用户冷启动自动兜底。", TEAL),
        ("C", "找得到", "3 公里内的真实门店",
         "GPS → 逆地理 → IP 兜底 → 默认位置四级定位；距离 / 评分 / 人均 / 品类四维筛选，\n地图标记与列表联动，一键导航。", "3B7DD8"),
        ("D", "跑得稳", "全链路可降级",
         "无高德 Key、无大模型、无外网、无索引时功能退化而不报错；零配置即可完成完整演示。", "8B5CF6"),
    ]
    for i, (letter, title, sub, body, accent) in enumerate(goals):
        col, row = i % 2, i // 2
        x = 0.62 + col * 6.12
        y = 1.6 + row * 1.99
        rect(s, x, y, 5.94, 1.84, WHITE, line=LINE, radius=0.06, shadow=True)
        rect(s, x, y, 0.06, 1.84, accent, shape=MSO_SHAPE.RECTANGLE)
        circle_num(s, x + 0.28, y + 0.24, 0.4, letter, fill=accent, size=13)
        tx(s, x + 0.86, y + 0.22, 2.0, 0.34, title, size=15, bold=True, color=INK)
        tx(s, x + 0.86, y + 0.58, 4.3, 0.3, sub, size=10, bold=True, color=accent)
        tx(s, x + 0.3, y + 0.95, 5.4, 0.82, body, size=9.5, color=GRAY, spacing=1.3)

    rect(s, 0.62, 5.62, 12.06, 0.95, CARD, radius=0.05)
    rich(s, 0.92, 5.78, 11.4, 0.7,
         [[("验收口径　", ORANGE_D, True, 10.5),
           ("62 道菜谱全部可检索　·　39 个 REST 接口可调　·　196 项测试全绿　·　断网仍有离线样例　·　首屏交互 < 2s",
            INK_SOFT, False, 10.5)]], spacing=1.3)


# ---------------------------------------------------------------- 03 架构
def slide_arch(prs, page):
    s = blank(prs)
    head(s, "03 / ARCHITECTURE", "系统架构：六层职责，一个 Django 项目",
         "单项目多 app，不引入 DRF；每层只依赖下一层的契约，任何外部能力都收口在能力层。", page)

    layers = [
        ("展示层", "Vue 3 全局构建（无打包步骤）· 暖色设计系统 · 单页壳 + 5 个业务视图 · SSE 流式渲染",
         "static/ templates/", "3C4663"),
        ("接口层", "39 个 REST 接口（25 + 14）· 统一 {code, msg, data} · 函数视图 + JsonResponse · 登录 / CSRF / 限流装饰器",
         "food/api_urls.py places/urls.py", ORANGE_D),
        ("业务层", "food：认证 · 菜谱 · 搜索 · 推荐 · 助手 · 后台　|　places：POI · 收藏 · 打卡 · Key 池 · 排序",
         "food/ places/", "2F6FB5"),
        ("算法层", "rag：中文切分 / hashing 向量 / BM25 / RRF　·　recommend：热度 / 画像 / item-CF / MMR　·　websearch：多引擎 + JSON-LD",
         "rag/ recommend/ websearch/", TEAL),
        ("能力层", "LLM 客户端（OpenAI 兼容，可换网关）· 高德 Web/JS Key 池 + 冷却调度 · 邮件服务 · 多引擎 HTTP 客户端",
         "RAGDjange/ places/keypool.py", "7C5CD6"),
        ("数据层", "PostgreSQL 18（可切 SQLite / MySQL）· 种子 JSON · 向量索引文件 · 离线 fixtures · 行为与搜索日志",
         "data/seed/ data/index/ fixtures/", "8A6A4F"),
    ]
    y = 1.55
    lh, gap = 0.72, 0.115
    for name, desc, path, color in layers:
        rect(s, 0.62, y, 1.42, lh, color, radius=0.05, shadow=True)
        tx(s, 0.62, y + 0.2, 1.42, 0.32, name, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        rect(s, 2.16, y, 8.55, lh, WHITE, line=LINE, radius=0.04, shadow=True)
        tx(s, 2.34, y + 0.08, 8.2, lh - 0.12, desc, size=9.3, color=INK_SOFT, spacing=1.28,
           anchor=MSO_ANCHOR.MIDDLE)
        tx(s, 10.86, y + 0.22, 2.0, 0.3, path, size=8.2, color=GRAY_L)
        y += lh + gap

    tx(s, 0.62, y + 0.06, 12.0, 0.3,
       "数据流：Vue 视图 → axios(统一拆包) → 接口层校验 → 服务层组装 → 算法层打分 → 外部能力(带缓存/降级) → PostgreSQL",
       size=9.5, color=GRAY)


# ---------------------------------------------------------------- 04 找菜
def slide_recipes(prs, page):
    s = blank(prs)
    head(s, "04 / FEATURES", "核心功能 · 主线一：在家做（找菜）",
         "从「有菜谱」到「懂你想吃什么」：数据 → 检索 → 推荐 → 对话 → 采集，形成闭环。", page)

    top = [
        ("01", "菜谱库（真实数据）",
         "62 道真实菜谱 · 12 菜系 · 34 标签 · 132 食材\n553 条食材用量、356 个制作步骤、热量与营养估算\n全部由 seed_data 幂等灌入，可重复重建"),
        ("02", "语义检索（自研 RAG）",
         "中文 1–3 gram 切分 + hashing TF-IDF(1024 维) 与 BM25 双路召回\nRRF 融合排序，命中块携带菜系/标签/时长元信息\n零模型依赖，可选接 Ollama embedding 升级"),
        ("03", "混合推荐（可解释）",
         "内容画像 / item-CF / 热度衰减 / 新鲜度四路加权\nMMR 去同质化，每条结果都带一句中文推荐理由\n冷启动无行为数据时自动切换热度 + 菜系多样性"),
    ]
    for i, (num, title, body) in enumerate(top):
        feature_card(s, 0.62 + i * 4.14, 1.62, 3.94, 2.28, num, title, body, body_size=9.5)

    bottom = [
        ("04", "AI 美食助手",
         "SSE 流式对话 + 引用溯源（点开即回到菜谱原文）\n5 类工具调用：检索 / 推荐 / 联网 / 统计 / 邮件\n模型不可用时降级为「检索式回答」，功能不中断"),
        ("05", "联网菜谱采集",
         "DDG / Bing / 360 多引擎搜索，失败自动切换下一个\n解析 schema.org Recipe（含 @graph 嵌套），统一结构\n按 source_url 去重入库，管理与命令双入口"),
    ]
    for i, (num, title, body) in enumerate(bottom):
        feature_card(s, 0.62 + i * 6.12, 4.06, 5.94, 1.95, num, title, body, body_size=9.5,
                     accent=TEAL, fill=CARD_ALT)


# ---------------------------------------------------------------- 05 找店
def slide_places(prs, page):
    s = blank(prs)
    head(s, "04 / FEATURES", "核心功能 · 主线二：出门吃（找店）",
         "接入高德 POI 的真实门店：定位 → 搜索 → 排序 → 地图 → 到店资产，全链路含配额与降级处理。", page)

    cards = [
        ("01", "四级定位", "浏览器 GPS → 高德逆地理（省市区/商圈）→ IP 定位 → 默认位置。\n任一级失败自动下沉，保证「总有数据可显示」。"),
        ("02", "周边搜索", "餐饮大类 050000，半径 500m–10km 可调；品类 / 关键词 / 人均预算筛选；\nPOI 结果 10 分钟缓存，不长期落库（避免过期脏数据）。"),
        ("03", "五路智能排序", "距离衰减 × 评分 × 人均预算匹配 × 口味画像 × 近期热度，五路加权且权重可配；\n排序单调性有单元测试守护。"),
        ("04", "地图与到店资产", "高德 JS API 2.0 标记点 + 信息窗体 + 自动视野；导航/叫车走 URI 不耗服务端配额；\n收藏店铺快照、到店打卡（评分/推荐菜/人均）并聚合均分。"),
    ]
    for i, (num, title, body) in enumerate(cards):
        col, row = i % 2, i // 2
        feature_card(s, 0.62 + col * 6.12, 1.6 + row * 1.72, 5.94, 1.58, num, title, body,
                     body_size=9.5, accent="3B7DD8", fill=WHITE)

    rect(s, 0.62, 5.22, 12.06, 1.42, INK, radius=0.05, shadow=True)
    tx(s, 0.94, 5.4, 11.4, 0.3, "Key 池与配额治理：把「不可用」做成一种正常状态", size=12.5, bold=True, color=WHITE)
    body = ("多 Key 轮询分发 · 配额类错误（10003 / 10029）池内冷却 6 小时 · QPS 类错误（10021）短冷却 · "
            "全池耗尽自动降级为内置离线样例并标记 degraded=true")
    tx(s, 0.94, 5.76, 11.5, 0.6, body, size=9.5, color="AEB6C7", spacing=1.35)
    tx(s, 0.94, 6.22, 11.5, 0.3, "前端仍能完整演示，只是数据来源从「实时」变为「样例」—— 降级对用户可见、可解释。",
       size=9.5, color="FFB37A")


# ---------------------------------------------------------------- 06 RAG 算法
def slide_rag(prs, page):
    s = blank(prs)
    head(s, "05 / ALGORITHM", "算法一：RAG 混合检索（向量 + 词法 + RRF）",
         "目标不是炫技：在「不下载任何模型」的前提下拿到可用的语义召回。实测 62 道菜 → 99 个检索块 / 1024 维 / 293ms。", page)

    steps = [
        ("① 切分", "chunker", "按段落/句子切 420 字、80 重叠\n优先不切断步骤\nchunk_id = recipe:12:0"),
        ("② 向量召回", "hashing TF-IDF", "char 1–3 gram + 词 bigram\nhashing → 1024 维，log 饱和 TF\nL2 归一化后余弦 top-N"),
        ("③ 词法召回", "BM25", "自实现 BM25（k1=1.5, b=0.75）\n中文 bigram 分词 + 英文数字词\n统计 df / idf / avgdl"),
        ("④ 融合", "RRF", "score = Σ 1/(60 + rank)\nHYBRID_BM25_WEIGHT 调权\n单路为空退化为另一路"),
    ]
    x = 0.62
    for i, (title, sub, body) in enumerate(steps):
        rect(s, x, 1.6, 2.62, 2.16, WHITE, line=LINE, radius=0.06, shadow=True)
        rect(s, x, 1.6, 2.62, 0.055, ORANGE, shape=MSO_SHAPE.RECTANGLE)
        tx(s, x + 0.22, 1.76, 2.2, 0.3, title, size=12, bold=True, color=INK)
        tx(s, x + 0.22, 2.06, 2.2, 0.26, sub, size=9, bold=True, color=ORANGE_D)
        tx(s, x + 0.22, 2.36, 2.24, 1.3, body, size=8.8, color=GRAY, spacing=1.32)
        if i < 3:
            arrow(s, x + 2.7, 2.58, 0.3, 0.18)
        x += 3.13

    # 底部三块
    blocks = [
        ("零依赖可交付", "默认 hashing 向量，不联网、不下载模型即可跑通；\n可选切换 Ollama embedding，维度变化由 manifest 熔断（不匹配即重建）。"),
        ("可解释", "每个 Hit 携带 title / cuisine / tags / 总时长，\n前端直接把「为什么是这道菜」展示成引用卡片。"),
        ("不炸上层", "search() 内任何异常都返回 [] 并记 warning；\n索引缺失时上层自动退化为关键词检索，页面不 500。"),
    ]
    for i, (title, body) in enumerate(blocks):
        x = 0.62 + i * 4.14
        rect(s, x, 4.06, 3.94, 1.42, CARD, radius=0.06)
        tx(s, x + 0.24, 4.2, 3.4, 0.3, title, size=11.5, bold=True, color=ORANGE_D)
        tx(s, x + 0.24, 4.52, 3.5, 0.9, body, size=9, color=GRAY, spacing=1.32)

    rect(s, 0.62, 5.68, 12.06, 0.9, CARD_ALT, radius=0.05, line=LINE)
    rich(s, 0.92, 5.82, 11.5, 0.6,
         [[("为什么两路一起用　", TEAL, True, 10),
           ("向量抓「意思相近但用词不同」（番茄 vs 西红柿），BM25 抓「精确词命中」（麻婆豆腐）；"
            "RRF 只看排名不看分数尺度，天然免疫两路分数量纲不一致的问题。", INK_SOFT, False, 9.8)]],
         spacing=1.35)


# ---------------------------------------------------------------- 07 推荐算法
def slide_recommend(prs, page):
    s = blank(prs)
    head(s, "05 / ALGORITHM", "算法二：四路混合推荐 + MMR 去同质化",
         "纯 numpy / dict 实现，不引入 scipy 与 sklearn；每一路都可单测、可解释、可单独关掉。", page)

    # 左侧权重图
    rect(s, 0.62, 1.6, 5.6, 3.5, WHITE, line=LINE, radius=0.06, shadow=True)
    tx(s, 0.88, 1.76, 4.6, 0.3, "默认权重（settings.RECOMMEND）", size=10, bold=True, color=GRAY)
    weights = [("内容画像", 0.45, ORANGE), ("协同过滤 item-CF", 0.30, TEAL),
               ("热度衰减", 0.15, "3B7DD8"), ("新鲜度", 0.10, "8B5CF6")]
    y = 2.16
    for name, val, color in weights:
        tx(s, 0.88, y, 2.3, 0.3, name, size=10, bold=True, color=INK_SOFT)
        rect(s, 2.9, y + 0.075, 2.4, 0.2, "F1ECE6", radius=0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(s, 2.9, y + 0.075, max(0.12, 2.4 * val / 0.45), 0.2, color, radius=0.5)
        tx(s, 5.4, y + 0.02, 0.7, 0.3, f"{val:.2f}", size=10.5, bold=True, color=color, font=FONT_NUM)
        y += 0.62
    tx(s, 0.88, 4.62, 4.9, 0.4, "四路归一化到 0–1 后加权求和，MMR 再按 λ 做批次级重排。",
       size=9, color=GRAY, spacing=1.3)

    # 右侧说明
    points = [
        ("内容画像", "菜系 / 标签权重、食材白名单、忌口、可用时长、辣度、饮食标签共同打分，同步产出中文理由（「你常吃川菜」）。"),
        ("item-CF", "Rating(1.0) + 行为加权（view 1 / click 1.5 / favorite 3 / cook_done 4）构建稀疏矩阵，进程内缓存相似表。"),
        ("热度衰减", "engagement / (hours + 2)^1.5，随时间自然降温；定时命令批量回写 hot_score，避免实时算。"),
        ("MMR", "λ 平衡相关性与多样性，解决「一屏全是川菜」；开箱可关，方便对比实验。"),
        ("冷启动", "无行为数据时切换热度 + 菜系多样性铺底，新用户首屏不会空。"),
    ]
    y = 1.6
    for name, body in points:
        rect(s, 6.5, y, 6.18, 0.9, CARD if y % 0.6 < 0.3 else CARD, radius=0.05)
        tx(s, 6.72, y + 0.12, 1.4, 0.3, name, size=10.5, bold=True, color=ORANGE_D)
        tx(s, 8.02, y + 0.1, 4.5, 0.76, body, size=9, color=GRAY, spacing=1.3)
        y += 0.94

    rich(s, 6.5, 6.34, 6.2, 0.42,
         [[("性能预算　", TEAL, True, 9.5),
           ("候选集 ≤ 500 → 全流程 < 200ms，缓存命中 < 50ms（本地 PostgreSQL 实测）", INK_SOFT, False, 9.5)]],
         spacing=1.2)


# ---------------------------------------------------------------- 08 工程质量
def slide_quality(prs, page):
    s = blank(prs)
    head(s, "06 / ENGINEERING", "工程质量：把「会挂」的地方全都设计成降级",
         "外部依赖（地图、大模型、外网、索引、数据库方言）全部收口，失败只降功能、不降可用性。", page)

    stats = [("196", "项", "pytest 全绿"), ("74", "个", "Python 模块"),
             ("15,752", "行", "Python 代码"), ("22,000", "行", "前端 JS/CSS")]
    for i, (v, u, label) in enumerate(stats):
        stat_card(s, 0.62 + i * 3.11, 1.58, 2.92, 1.42, v, u, label, accent=ORANGE if i < 2 else TEAL)

    rows = [
        ["外部能力", "缺失 / 异常时的行为", "对用户的影响"],
        ["高德 Key 缺失或配额耗尽", "返回内置离线 POI 样例，标记 degraded=true", "附近页照常可用，标注数据来源"],
        ["LLM 不可用 / 超时", "助手降级为检索式回答，引用照常返回", "回答变简洁，结论仍可溯源"],
        ["RAG 索引未构建", "语义检索退化为关键词检索，不抛 500", "搜索仍可用，精度略降"],
        ["外网断开", "websearch 读取本地 fixtures 并标记", "演示链路不断"],
        ["数据库方言差异", "PG 守卫测试：禁裸 SQL、禁硬编码 DSN", "换库不改代码"],
    ]
    table(s, 0.62, 3.18, [3.3, 4.7, 4.06], rows, row_h=0.49, header_h=0.38, body_size=9.3, head_size=10)

    rect(s, 0.62, 6.5, 12.06, 0.38, CARD, radius=0.04)
    rich(s, 0.86, 6.56, 11.6, 0.28,
         [[("契约驱动并行　", ORANGE_D, True, 9.5),
           ("docs/CONTRACTS.md 固定全部函数签名 + 目录级写权限表，6 人同时改代码零冲突；序列化只走 Recipe.to_card() / to_detail() 单一出口。",
            INK_SOFT, False, 9.5)]], spacing=1.15)


# ---------------------------------------------------------------- 09 数据看板
def slide_data(prs, page):
    s = blank(prs)
    head(s, "07 / DATA", "数据规模看板：可复现的演示底座",
         "所有数字来自当前运行实例（PostgreSQL 实查），种子数据幂等，任何机器上都能一键重建。", page)

    cards = [
        ("62", "道", "真实菜谱", "含用量、步骤、营养"),
        ("12", "个", "菜系分类", "川 / 粤 / 鲁 / 苏 / 浙…"),
        ("34", "个", "口味与场景标签", "5 类 kind，可多选"),
        ("132", "种", "食材", "带热量与营养素"),
        ("553", "条", "食材用量", "RecipeIngredient"),
        ("356", "步", "制作步骤", "含单步时长"),
        ("39", "个", "REST 接口", "food 25 + places 14"),
        ("196", "项", "自动化测试", "含降级与并发"),
    ]
    for i, (v, u, label, note) in enumerate(cards):
        col, row = i % 4, i // 4
        stat_card(s, 0.62 + col * 3.11, 1.6 + row * 1.9, 2.92, 1.7, v, u, label, note,
                  accent=ORANGE if row == 0 else TEAL)

    rect(s, 0.62, 5.5, 12.06, 1.28, INK, radius=0.05)
    tx(s, 0.92, 5.66, 5.0, 0.3, "行为数据与检索日志", size=11, bold=True, color="FFB37A")
    rich(s, 0.92, 5.98, 11.4, 0.7,
         [[("143 条评分　·　208 条行为埋点　·　61 次收藏　·　搜索/对话日志全量记录",
            WHITE, True, 11),
           ("　—　推荐与协同过滤的输入就来自这些真实行为，而不是造出来的假分数。", "9AA1AE", False, 9.5)]],
         spacing=1.3)


# ---------------------------------------------------------------- 10 界面
def slide_ui(prs, page):
    s = blank(prs)
    head(s, "07 / DEMO", "界面实拍：三条主路径",
         "登录 → 发现（推荐流）→ 附近（地图 + 门店列表），均为本地实例真实截图。", page)

    shots = [
        ("ui-login.png", "登录页", "邮箱验证码登录，游客模式一键体验；左侧品牌区讲清三个卖点。", 1.732),
        ("ui-home.png", "发现页 · 推荐流", "自然语言搜索 + 场景标签 + 大家都在吃 + 按菜系逛 + 本周热门。", 1.732),
        ("ui-nearby.png", "附近页 · 门店列表 + 地图", "四级定位到商圈，25 家 3km 内门店，左侧列表右侧地图联动。", 1.732),
    ]
    x = 0.62
    for fname, title, desc, ratio in shots:
        path = ASSETS / fname
        rect(s, x, 1.56, 3.94, 3.36, WHITE, line=LINE, radius=0.05, shadow=True)
        pic = s.shapes.add_picture(str(path), I(x + 0.08), I(1.64), width=I(3.78))
        h = 3.78 / ratio
        # 图片按宽度等比，若超卡片内高再裁掉底部
        if 1.64 + h > 4.84:
            pic.crop_bottom = (h - (4.84 - 1.64)) / h
        tx(s, x + 0.18, 4.98, 3.6, 0.32, title, size=11.5, bold=True, color=INK)
        tx(s, x + 0.18, 5.3, 3.66, 0.5, desc, size=8.8, color=GRAY, spacing=1.26)
        x += 4.14

    rect(s, 0.62, 5.92, 12.06, 0.86, CARD, radius=0.05)
    rich(s, 0.92, 6.06, 11.5, 0.62,
         [[("演示脚本　", ORANGE_D, True, 10),
           ("登录 → 语义搜「20 分钟快手菜」→ 打开详情看推荐理由与相似菜 → 切「附近」看地图与五路排序 → "
            "收藏 + 打卡 → 问助手「附近吃什么」，它直接调 POI 工具回答。", INK_SOFT, False, 9.5)]],
         spacing=1.35)


# ---------------------------------------------------------------- 11 分工
def slide_team(prs, page):
    s = blank(prs)
    head(s, "08 / TEAM", "团队分工：模块级负责，契约级协作",
         "6 人并行开发同一仓库：以 CONTRACTS.md 为唯一约定，按目录划分写权限，接口一次对齐。", page)

    members = [
        ("孔哲昊", "项目组长 · 架构与集成",
         "Django 配置与总路由、统一 API 约定、页面视图、LLM 客户端、契约维护、代码评审与合并、整体进度",
         True),
        ("彭冰钦", "RAG 检索", "chunker 中文切分、hashing TF-IDF 向量、BM25 索引、向量库持久化、RRF 混合检索与单测", False),
        ("姚一飞", "推荐引擎", "热度时间衰减、内容画像与推荐理由、item-CF 相似表、混合排序与 MMR 去同质化", False),
        ("贺江涛", "附近美食", "高德客户端、Key 池与冷却、WGS-84/GCJ-02 坐标转换、POI 缓存、五路排序与降级", False),
        ("倪永硕", "联网采集", "多引擎搜索、schema.org Recipe 解析、入库去重、离线 fixtures、断网降级链路", False),
        ("易子涵", "前端与数据", "Vue 组件与暖色设计系统、页面交互联调、地图与列表联动、种子数据与演示脚本", False),
    ]
    for i, (name, role, work, lead) in enumerate(members):
        col, row = i % 3, i // 3
        x = 0.62 + col * 4.14
        y = 1.6 + row * 2.1
        fill = ORANGE if lead else WHITE
        rect(s, x, y, 3.94, 1.92, fill, line=None if lead else LINE, radius=0.06, shadow=True)
        tx(s, x + 0.26, y + 0.22, 3.2, 0.36, name, size=15, bold=True, color=WHITE if lead else INK)
        tx(s, x + 0.26, y + 0.62, 3.4, 0.3, role, size=9.6, bold=True,
           color="FFE3CC" if lead else ORANGE_D)
        tx(s, x + 0.26, y + 0.94, 3.44, 0.9, work, size=8.8,
           color="FFEDE0" if lead else GRAY, spacing=1.32)


# ---------------------------------------------------------------- 12 难点
def slide_challenges(prs, page):
    s = blank(prs)
    head(s, "08 / REVIEW", "项目难点与对策：五个真实踩过的坑",
         "每个难点都有对应的工程动作，而不是「加强沟通、继续努力」这类空话。", page)

    rows = [
        ["难点", "对策", "结果"],
        ["外部依赖全部不可控\n（Key / 模型 / 外网）", "分层降级 + 离线 fixtures + Key 池冷却调度", "零配置、断网也能完成完整演示"],
        ["6 人并行改同一仓库", "契约先行（CONTRACTS.md）+ 目录级写权限表", "合并零冲突，接口一次对齐"],
        ["推荐同质化与冷启动", "MMR 重排 + 内容画像兜底 + 热度铺底", "新用户首屏即有可解释推荐"],
        ["本地 14B 模型首 token 慢", "SSE 流式输出 + 工具调用替代长上下文", "首字可见时间显著下降，交互不阻塞"],
        ["坐标系混用导致标记偏移", "入参 WGS-84 自动转 GCJ-02，响应 center 恒为 GCJ-02", "地图标记与列表位置零偏移"],
    ]
    table(s, 0.62, 1.6, [3.5, 4.5, 4.06], rows, row_h=0.78, header_h=0.4, body_size=9.2, head_size=10)

    rect(s, 0.62, 6.42, 12.06, 0.42, CARD_ALT, radius=0.04, line=LINE)
    rich(s, 0.88, 6.5, 11.5, 0.28,
         [[("一句话总结　", TEAL, True, 9.5),
           ("我们花的力气，八成不在「把功能做出来」，而在「让功能在什么都缺的时候还能用」。", INK_SOFT, False, 9.5)]],
         spacing=1.15)


# ---------------------------------------------------------------- 13 展望
def slide_next(prs, page):
    s = blank(prs)
    head(s, "08 / NEXT", "不足与后续规划",
         "诚实地讲清现状：哪些是真做完了，哪些只是「接口留好了」。", page)

    items = [
        ("索引与评估", "上线前固化 build_index 产物；引入 NDCG@10、召回率离线评估，配合埋点做 A/B。", ORANGE),
        ("数据资产", "POI 热点结果冷存 + 异步刷新，沉淀「店—菜」关联，支撑更细的推荐。", TEAL),
        ("体验", "移动端 PWA 与小程序；地图聚合、路线规划、多人拼单场景。", "3B7DD8"),
        ("风险治理", "高德配额与合规边界；模型幻觉用「引用强制溯源」兜底，无来源不外发。", "8B5CF6"),
    ]
    y = 1.6
    for title, body, color in items:
        rect(s, 0.62, y, 7.5, 1.16, WHITE, line=LINE, radius=0.05, shadow=True)
        rect(s, 0.62, y, 0.055, 1.16, color, shape=MSO_SHAPE.RECTANGLE)
        tx(s, 0.92, y + 0.16, 2.2, 0.32, title, size=12, bold=True, color=INK)
        tx(s, 0.92, y + 0.52, 7.0, 0.6, body, size=9.5, color=GRAY, spacing=1.3)
        y += 1.26

    rect(s, 8.38, 1.6, 4.3, 4.9, CARD, radius=0.05)
    tx(s, 8.66, 1.8, 3.6, 0.32, "现状自评（不美化）", size=10, bold=True, color=ORANGE_D)
    notes = [
        ("已完成", "菜谱 / 推荐 / RAG / 助手 / 附近门店 / 收藏打卡，全链路可跑，196 项测试全绿。"),
        ("已留口", "RAG 索引与 embedding 可切换；POI 未做长期落库；单 JS Key 不做自动轮换。"),
        ("待补", "离线评测集、A/B 实验框架、移动端适配、运维监控面板。"),
    ]
    y = 2.2
    for tag, body in notes:
        tx(s, 8.66, y, 1.0, 0.28, tag, size=9.5, bold=True, color=TEAL if tag == "已完成" else GRAY)
        tx(s, 8.66, y + 0.3, 3.7, 1.0, body, size=9.2, color=INK_SOFT, spacing=1.34)
        y += 1.4


# ---------------------------------------------------------------- 14 总结
def slide_end(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, INK, shape=MSO_SHAPE.RECTANGLE)
    c1 = s.shapes.add_shape(MSO_SHAPE.OVAL, I(9.4), I(-1.9), I(4.6), I(4.6))
    c1.fill.solid()
    c1.fill.fore_color.rgb = rgb("242A3D")
    c1.line.fill.background()
    c1.shadow.inherit = False
    c2 = s.shapes.add_shape(MSO_SHAPE.OVAL, I(-1.3), I(4.9), I(3.2), I(3.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = rgb(ORANGE)
    c2.line.fill.background()
    c2.shadow.inherit = False

    tx(s, 0.92, 1.1, 10.0, 0.34, "THANKS FOR LISTENING", size=10.5, bold=True, color="FFB37A")
    tx(s, 0.92, 1.5, 11.2, 1.9,
       "把「今天吃什么」\n从一道选择题，做成一句话的检索。", size=32, bold=True, color=WHITE, spacing=1.3)
    rect(s, 0.95, 3.44, 1.18, 0.075, ORANGE, shape=MSO_SHAPE.RECTANGLE)

    gains = [("01", "契约先行的多人协作", "签名先定、权限先分，再并行写码"),
             ("02", "零依赖的检索与推荐", "不靠模型也能有语义检索与个性化"),
             ("03", "可降级的工程习惯", "把失败当常态设计，而不是当事故处理")]
    x = 0.95
    for num, title, body in gains:
        tx(s, x, 3.86, 0.7, 0.32, num, size=13, bold=True, color=ORANGE, font=FONT_NUM)
        tx(s, x + 0.5, 3.84, 3.4, 0.34, title, size=11.5, bold=True, color=WHITE)
        tx(s, x + 0.5, 4.2, 3.35, 0.6, body, size=9, color="9AA1AE", spacing=1.3)
        x += 4.1

    rect(s, 0.95, 5.3, 11.4, 0.012, "3A4257", shape=MSO_SHAPE.RECTANGLE)
    tx(s, 0.95, 5.52, 5.6, 0.5, "第十组　·　组长 孔哲昊", size=13, bold=True, color=WHITE)
    tx(s, 0.95, 5.9, 7.0, 0.4, "彭冰钦　姚一飞　贺江涛　倪永硕　易子涵", size=11, color="AEB6C7")
    pill(s, 9.9, 5.55, 2.45, 0.62, "Q & A", fill=ORANGE, color=WHITE, size=16)


# ---------------------------------------------------------------- 主流程
def main():
    prs = Presentation()
    prs.slide_width = I(W)
    prs.slide_height = I(H)

    slide_cover(prs)
    slide_toc(prs, 2)
    slide_background(prs, 3)
    slide_goals(prs, 4)
    slide_arch(prs, 5)
    slide_recipes(prs, 6)
    slide_places(prs, 7)
    slide_rag(prs, 8)
    slide_recommend(prs, 9)
    slide_quality(prs, 10)
    slide_data(prs, 11)
    slide_ui(prs, 12)
    slide_team(prs, 13)
    slide_challenges(prs, 14)
    slide_next(prs, 15)
    slide_end(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"OK -> {OUT}  slides={len(prs.slides.__iter__.__self__._sldIdLst)}")


if __name__ == "__main__":
    sys.exit(main())
