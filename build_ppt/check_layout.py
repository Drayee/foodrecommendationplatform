#!/usr/bin/env python3
"""用真实字体度量估算每个文本框的排版高度，找出溢出文本框的形状。"""
from __future__ import annotations

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

DPI = 96.0
FONT_REG = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"
_cache: dict = {}


def font_for(size_pt: float, bold: bool):
    key = (round(size_pt, 1), bold)
    if key not in _cache:
        path = FONT_BOLD if bold else FONT_REG
        _cache[key] = ImageFont.truetype(path, max(6, int(round(size_pt * DPI / 72.0))), index=0)
    return _cache[key]


def text_px_width(text: str, size_pt: float, bold: bool) -> float:
    f = font_for(size_pt, bold)
    return f.getlength(text)


def para_lines(runs, limit_pt: float) -> int:
    """贪心按字符换行，返回行数。limit_pt 为可用宽度（pt）。"""
    if not runs:
        return 1
    # 合并成 (char, size_pt, bold)
    chars = []
    for text, size, bold in runs:
        for ch in text:
            chars.append((ch, size, bold))
    lines, cur = 1, 0.0
    for ch, size, bold in chars:
        w = text_px_width(ch, size, bold) * 72.0 / DPI
        if ch == "\n":
            lines += 1
            cur = 0.0
            continue
        if cur + w > limit_pt and cur > 0:
            lines += 1
            cur = w
        else:
            cur += w
    return lines


def analyze(pptx_path: Path):
    prs = Presentation(str(pptx_path))
    problems = []
    for idx, slide in enumerate(prs.slides, start=1):
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            text = shp.text_frame.text.strip()
            if not text:
                continue
            tf = shp.text_frame
            w_pt = (shp.width - (tf.margin_left or 0) - (tf.margin_right or 0)) / 12700.0
            h_pt = (shp.height - (tf.margin_top or 0) - (tf.margin_bottom or 0)) / 12700.0
            if w_pt <= 0 or h_pt <= 0:
                continue
            need = 0.0
            for p in tf.paragraphs:
                runs = []
                max_size = 12.0
                for r in p.runs:
                    size = r.font.size.pt if r.font.size else 12.0
                    max_size = max(max_size, size)
                    runs.append((r.text, size, bool(r.font.bold)))
                spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
                n = para_lines(runs, w_pt)
                need += n * max_size * 1.22 * spacing
                if p.space_after:
                    need += p.space_after.pt
            if need > h_pt * 1.04:
                problems.append((idx, round(need, 1), round(h_pt, 1), text[:46].replace("\n", " / ")))
    return problems


def main():
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    if target is None:
        print("usage: check_layout.py <pptx>")
        return 2
    probs = analyze(target)
    if not probs:
        print("no overflow detected")
    for page, need, have, preview in probs:
        print(f"P{page:02d} need={need}pt have={have}pt  +{need - have:.0f}pt  | {preview}")
    print(f"\ntotal suspects: {len(probs)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
